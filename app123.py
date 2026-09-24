# -*- coding: utf-8 -*-
"""
ERP Bank Reconciliation Tool
============================
Streamlit application that reconciles an ERP export against a bank statement (RLV)
in three sequential phases:

    Phase 1 : ERP.reference found inside RLV.description      + same amount
    Phase 2 : fuzzy(ERP.description, RLV.description) >= 80%  + same amount + |date gap| <= 15 d
    Phase 3 : fuzzy(ERP.nom,         RLV.description) >= 80%  + same amount + |date gap| <= 15 d

A row matched in one phase is never reused. Columns are auto-detected even when
their names vary slightly.

Run:
    pip install streamlit pandas openpyxl rapidfuzz
    streamlit run app.py
"""

from __future__ import annotations

import io
import re
import unicodedata
import datetime
from collections import defaultdict

import pandas as pd
import streamlit as st
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from rapidfuzz import fuzz


# ============================================================================
#  Configuration
# ============================================================================
FUZZY_THRESHOLD = 80        # RapidFuzz minimum score (%) for Phase 2 & 3
MAX_DATE_GAP_DAYS = 15      # maximum |ERP date - RLV date| for Phase 2 & 3
REPORT_NAME = "Reconciliation_Final_Report.xlsx"

# Keyword dictionary used for fuzzy column detection (accent/case-insensitive).
ROLE_KEYWORDS = {
    "date":        ["date"],
    "credit":      ["credit", "cred", "avoir"],
    "debit":       ["debit", "deb"],
    "document":    ["document", "piece", "justif", "ndoc", "n doc", "num piece"],
    "reference":   ["reference", "ref"],
    "nom":         ["nom", "client", "tiers", "beneficiaire", "raison", "fournisseur", "donneur"],
    "description": ["description", "libelle", "libell", "operation", "detail",
                    "motif", "narration", "objet", "intitule", "desc"],
}


# ============================================================================
#  Generic helpers
# ============================================================================
def deaccent(value) -> str:
    """Lower-case, accent-stripped, trimmed string."""
    s = unicodedata.normalize("NFD", str(value if value is not None else ""))
    return "".join(c for c in s if unicodedata.category(c) != "Mn").strip().lower()


def clean_text(value) -> str:
    s = str(value).strip()
    return "" if s.lower() in ("nan", "none", "nat") else s


def to_number(series: pd.Series) -> pd.Series:
    """Parse a money column tolerating spaces and comma decimals ('1 234,56')."""
    if pd.api.types.is_numeric_dtype(series):
        return pd.to_numeric(series, errors="coerce").fillna(0.0)
    cleaned = (series.astype(str)
               .str.replace(r"\s", "", regex=True)
               .str.replace(" ", "", regex=False)
               .str.replace(",", ".", regex=False)
               .str.replace(r"[^0-9.\-]", "", regex=True))
    return pd.to_numeric(cleaned, errors="coerce").fillna(0.0)


def parse_date(value):
    """Robust per-value date parser: real dates, dd/mm/yyyy, dd-mm-yyyy, yyyy-mm-dd."""
    if isinstance(value, (pd.Timestamp, datetime.datetime, datetime.date)):
        return pd.Timestamp(value).normalize()
    s = str(value).strip()
    if not s:
        return pd.NaT
    parts = re.split(r"[-/.]", s.split()[0])
    if len(parts) == 3 and all(p.isdigit() for p in parts):
        a, b, c = (int(p) for p in parts)
        try:
            return pd.Timestamp(year=a, month=b, day=c) if a > 31 else pd.Timestamp(year=c, month=b, day=a)
        except ValueError:
            return pd.NaT
    return pd.to_datetime(s, dayfirst=True, errors="coerce")


# ============================================================================
#  Loading & column detection
# ============================================================================
def read_excel(upload) -> pd.DataFrame:
    """Read an uploaded Excel file; prefer an ERP/RLV-named sheet, else the first."""
    sheets = pd.read_excel(upload, sheet_name=None)
    for name in sheets:
        if deaccent(name) in ("erp", "rlv", "releve", "bank", "banque"):
            return sheets[name]
    return list(sheets.values())[0]


def detect_columns(df: pd.DataFrame) -> dict:
    """Map each logical role to an actual column name (or None)."""
    headers = {c: deaccent(c) for c in df.columns}
    used, mapping = set(), {}
    for role, keywords in ROLE_KEYWORDS.items():
        chosen = None
        for col in df.columns:
            if col in used:
                continue
            if any(kw in headers[col] for kw in keywords):
                chosen = col
                break
        mapping[role] = chosen
        if chosen is not None:
            used.add(chosen)
    return mapping


def normalize(df: pd.DataFrame) -> tuple[pd.DataFrame, dict]:
    """Return a normalized working frame + the detected column mapping."""
    cmap = detect_columns(df)
    out = pd.DataFrame(index=df.index)

    out["date"] = df[cmap["date"]].map(parse_date) if cmap["date"] else pd.NaT

    credit = to_number(df[cmap["credit"]]).abs() if cmap["credit"] else pd.Series(0.0, index=df.index)
    debit = to_number(df[cmap["debit"]]).abs() if cmap["debit"] else pd.Series(0.0, index=df.index)
    out["debit"] = debit
    out["credit"] = credit
    out["amount"] = credit.where(credit > 0, debit)                 # transaction value
    out["cents"] = (out["amount"] * 100).round().astype("int64")

    out["reference"] = df[cmap["reference"]].map(clean_text) if cmap["reference"] else ""
    out["nom"] = df[cmap["nom"]].map(clean_text) if cmap["nom"] else ""
    out["document"] = df[cmap["document"]].map(clean_text) if cmap["document"] else ""

    if cmap["description"]:
        out["description"] = df[cmap["description"]].map(clean_text)
    else:  # fall back to any leftover text column(s)
        skip = {cmap[r] for r in cmap} | {cmap["date"]}
        text_cols = [c for c in df.columns if c not in skip
                     and not pd.api.types.is_numeric_dtype(df[c])
                     and not pd.api.types.is_datetime64_any_dtype(df[c])]
        out["description"] = (df[text_cols].astype(str).agg(" ".join, axis=1)
                              if text_cols else "")

    out["ref_norm"] = out["reference"].map(deaccent)
    out["nom_norm"] = out["nom"].map(deaccent)
    out["desc_norm"] = out["description"].map(deaccent)
    return out, cmap


# ============================================================================
#  Reconciliation engine
# ============================================================================
def reconcile(erp: pd.DataFrame, rlv: pd.DataFrame) -> list[dict]:
    """Sequential 3-phase matching. Returns a list of match dicts."""
    rlv_by_amt: dict[int, list] = defaultdict(list)
    for rj in rlv.index:
        if rlv.at[rj, "cents"] > 0:
            rlv_by_amt[rlv.at[rj, "cents"]].append(rj)

    erp_done, rlv_done, matches = set(), set(), []

    def date_gap_ok(ed, rd):
        return not pd.isna(ed) and not pd.isna(rd) and abs((ed - rd).days) <= MAX_DATE_GAP_DAYS

    # ---------- Phase 1 : reference inside description, same amount ----------
    for ei in erp.index:
        ref = erp.at[ei, "ref_norm"]
        cents = erp.at[ei, "cents"]
        if not ref or cents <= 0:
            continue
        for rj in rlv_by_amt.get(cents, []):
            if rj not in rlv_done and ref in rlv.at[rj, "desc_norm"]:
                erp_done.add(ei)
                rlv_done.add(rj)
                matches.append({"phase": "Phase 1", "score": 100.0, "ei": ei, "rj": rj})
                break

    # ---------- Phase 2 & 3 : fuzzy match, same amount, date gap <= 15 d ----------
    def fuzzy_phase(label, field):
        for ei in erp.index:
            if ei in erp_done:
                continue
            text = erp.at[ei, field]
            cents = erp.at[ei, "cents"]
            ed = erp.at[ei, "date"]
            if not text or cents <= 0:
                continue
            best_rj, best_score = None, FUZZY_THRESHOLD - 1
            for rj in rlv_by_amt.get(cents, []):
                if rj in rlv_done or not date_gap_ok(ed, rlv.at[rj, "date"]):
                    continue
                score = fuzz.token_set_ratio(text, rlv.at[rj, "desc_norm"])
                if score >= FUZZY_THRESHOLD and score > best_score:
                    best_score, best_rj = score, rj
            if best_rj is not None:
                erp_done.add(ei)
                rlv_done.add(best_rj)
                matches.append({"phase": label, "score": float(best_score), "ei": ei, "rj": best_rj})

    fuzzy_phase("Phase 2", "desc_norm")
    fuzzy_phase("Phase 3", "nom_norm")
    return matches


# ============================================================================
#  Report generation
# ============================================================================
def style_sheet(ws):
    navy, white = "1F3864", "FFFFFF"
    for c in range(1, ws.max_column + 1):
        cell = ws.cell(1, c)
        cell.font = Font(name="Arial", size=10, bold=True, color=white)
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[cell.column_letter].width = 20
    for col in ws.iter_cols(min_row=2):
        for cell in col:
            if isinstance(cell.value, (int, float)):
                cell.number_format = "#,##0.00"
    ws.freeze_panes = "A2"


def build_report(erp_raw, rlv_raw, erp, rlv, matches):
    erp_done = {m["ei"] for m in matches}
    rlv_done = {m["rj"] for m in matches}

    matched_rows = []
    for m in matches:
        e, r = erp.loc[m["ei"]], rlv.loc[m["rj"]]
        matched_rows.append({
            "Match Type": m["phase"], "Score (%)": round(m["score"], 1), "Montant": e["amount"],
            "ERP Date": e["date"], "ERP Référence": e["reference"], "ERP Nom": e["nom"],
            "ERP Document": e["document"], "ERP Description": e["description"],
            "ERP Débit": e["debit"], "ERP Crédit": e["credit"],
            "RLV Date": r["date"], "RLV Description": r["description"],
            "RLV Débit": r["debit"], "RLV Crédit": r["credit"],
        })
    matched_df = pd.DataFrame(matched_rows, columns=[
        "Match Type", "Score (%)", "Montant", "ERP Date", "ERP Référence", "ERP Nom",
        "ERP Document", "ERP Description", "ERP Débit", "ERP Crédit",
        "RLV Date", "RLV Description", "RLV Débit", "RLV Crédit"])

    unmatched_erp = erp_raw.loc[[i for i in erp.index if i not in erp_done]].reset_index(drop=True)
    unmatched_rlv = rlv_raw.loc[[i for i in rlv.index if i not in rlv_done]].reset_index(drop=True)

    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl", datetime_format="dd/mm/yyyy") as xl:
        matched_df.to_excel(xl, sheet_name="Matched_Results", index=False)
        unmatched_erp.to_excel(xl, sheet_name="Unmatched_ERP", index=False)
        unmatched_rlv.to_excel(xl, sheet_name="Unmatched_RLV", index=False)
    buf.seek(0)
    wb = load_workbook(buf)
    for ws in wb.worksheets:
        if ws.max_row >= 1:
            style_sheet(ws)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue(), matched_df, unmatched_erp, unmatched_rlv


# ============================================================================
#  Streamlit UI
# ============================================================================
def main():
    st.set_page_config(page_title="ERP Bank Reconciliation Tool", page_icon="🏦", layout="wide")
    st.title("🏦 ERP Bank Reconciliation Tool")
    st.caption("Rapprochement séquentiel en 3 phases (référence • description floue • nom flou) "
               "avec détection automatique des colonnes.")

    c1, c2 = st.columns(2)
    with c1:
        st.subheader("1️⃣ Fichier ERP")
        erp_file = st.file_uploader("ERP (.xlsx)", type=["xlsx", "xlsm", "xls"], key="erp")
    with c2:
        st.subheader("2️⃣ Relevé bancaire (RLV)")
        rlv_file = st.file_uploader("RLV (.xlsx)", type=["xlsx", "xlsm", "xls"], key="rlv")

    if not (erp_file and rlv_file):
        st.info("⬆️ Téléversez les deux fichiers (ERP et RLV) pour lancer le rapprochement.")
        return

    try:
        with st.spinner("Rapprochement en cours…"):
            erp_raw, rlv_raw = read_excel(erp_file), read_excel(rlv_file)
            erp, _ = normalize(erp_raw)
            rlv, _ = normalize(rlv_raw)
            matches = reconcile(erp, rlv)
            report_bytes, matched_df, un_erp, un_rlv = build_report(erp_raw, rlv_raw, erp, rlv, matches)
    except Exception as exc:
        st.error(f"Impossible de traiter les fichiers : {exc}")
        return

    p1 = sum(m["phase"] == "Phase 1" for m in matches)
    p2 = sum(m["phase"] == "Phase 2" for m in matches)
    p3 = sum(m["phase"] == "Phase 3" for m in matches)
    total = len(matches)

    st.subheader("📊 Indicateurs")
    r1 = st.columns(4)
    r1[0].metric("ERP Rows", len(erp_raw))
    r1[1].metric("RLV Rows", len(rlv_raw))
    r1[2].metric("Phase 1 Matches", p1)
    r1[3].metric("Phase 2 Matches", p2)
    r2 = st.columns(4)
    r2[0].metric("Phase 3 Matches", p3)
    r2[1].metric("Total Matches", total)
    r2[2].metric("Unmatched ERP", len(un_erp))
    r2[3].metric("Unmatched RLV", len(un_rlv))

    st.subheader("✅ Matched_Results")
    st.dataframe(matched_df, use_container_width=True, hide_index=True)
    with st.expander(f"⚠️ Unmatched ERP ({len(un_erp)})"):
        st.dataframe(un_erp, use_container_width=True, hide_index=True)
    with st.expander(f"⚠️ Unmatched RLV ({len(un_rlv)})"):
        st.dataframe(un_rlv, use_container_width=True, hide_index=True)

    st.download_button(
        "⬇️ Télécharger Reconciliation_Final_Report.xlsx",
        data=report_bytes,
        file_name=REPORT_NAME,
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        type="primary",
    )


if __name__ == "__main__":
    main()
